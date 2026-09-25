from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import os
import re
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, datetime
from html.parser import HTMLParser
from pathlib import Path

from app.core.config import settings
from app.services.knowledge.markdown_parser import ParsedMarkdown, parse_markdown

SUPPORTED_SUFFIXES = {".md", ".txt", ".pdf", ".docx", ".html", ".htm", ".json", ".csv"}
IGNORED_DIRS = {
    ".obsidian",
    ".git",
    "node_modules",
    "venv",
    ".venv",
    "__pycache__",
    "tmp",
    "cache",
}
IGNORED_FILES = {".raventech-knowledge-source.json"}
MAX_FILES_PER_SYNC = 20_000
MAX_EXTRACTED_TEXT_CHARS = 5_000_000
MAX_DOCX_UNCOMPRESSED_BYTES = 32 * 1024 * 1024
_SECRET_PATTERNS = (
    re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----", re.I),
    re.compile(
        r"(?i)\b(?:api[_-]?key|access[_-]?token|client[_-]?secret|password)\s*[:=]\s*[^\s]{8,}"
    ),
    re.compile(r"(?i)\b(?:AKIA[0-9A-Z]{16}|gh[pousr]_[A-Za-z0-9]{20,})\b"),
)


class DocumentParseError(ValueError):
    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class ParsedFile:
    relative_name: str
    content_type: str
    content: str
    file_hash: str
    size_bytes: int
    modified_at: datetime
    modified_ns: int
    parsed_markdown: ParsedMarkdown | None
    sensitive_content_warning: bool
    embed_hashes: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class ScanResult:
    files: list[ParsedFile]
    unsupported: int
    oversized: int
    unsafe_links: int
    scanned: int
    failed: int
    errors: list[dict[str, str]]
    unchanged: list[tuple[str, int, int, str]] = field(default_factory=list)


def scan_directory(
    root_path: str,
    *,
    max_file_bytes: int | None = None,
    known_hashes: dict[str, str] | None = None,
    force_parse_names: set[str] | None = None,
) -> ScanResult:
    """Read supported files below an attached root without following links."""
    root = Path(root_path).expanduser()
    if _is_link(root):
        raise DocumentParseError("source_root_is_link")
    try:
        canonical_root = root.resolve(strict=True)
    except OSError as exc:
        raise DocumentParseError("source_offline") from exc
    if not canonical_root.is_dir():
        raise DocumentParseError("source_root_not_directory")
    cap = max_file_bytes or settings.KNOWLEDGE_MAX_FILE_BYTES
    files: list[ParsedFile] = []
    unsupported = oversized = unsafe = scanned = failed = 0
    total_supported_bytes = 0
    errors: list[dict[str, str]] = []
    unchanged: list[tuple[str, int, int, str]] = []
    force_parse = force_parse_names or set()

    def note_walk_error(error: OSError) -> None:
        nonlocal failed
        failed += 1
        error_path = Path(error.filename) if error.filename else canonical_root
        try:
            relative = error_path.relative_to(canonical_root).as_posix()
        except ValueError:
            relative = ""
        errors.append({"relative_name": relative, "code": "directory_read_failed"})

    for current, dirnames, filenames in os.walk(
        canonical_root, topdown=True, onerror=note_walk_error, followlinks=False
    ):
        current_path = Path(current)
        safe_dirs: list[str] = []
        for name in dirnames:
            child = current_path / name
            if name.casefold() in IGNORED_DIRS:
                continue
            if _is_link(child):
                unsafe += 1
                errors.append(
                    {
                        "relative_name": child.relative_to(canonical_root).as_posix(),
                        "code": "unsafe_directory_link",
                    }
                )
                continue
            try:
                if child.resolve(strict=True).is_relative_to(canonical_root):
                    safe_dirs.append(name)
            except (OSError, ValueError):
                unsafe += 1
        dirnames[:] = sorted(safe_dirs, key=lambda value: (value.casefold(), value))
        for name in sorted(filenames, key=lambda value: (value.casefold(), value)):
            if name.casefold() in IGNORED_FILES:
                continue
            scanned += 1
            if scanned > MAX_FILES_PER_SYNC:
                raise DocumentParseError("source_file_count_limit")
            path = current_path / name
            if _is_link(path):
                unsafe += 1
                errors.append(
                    {
                        "relative_name": path.relative_to(canonical_root).as_posix(),
                        "code": "unsafe_link",
                    }
                )
                continue
            try:
                resolved = path.resolve(strict=True)
                if not resolved.is_relative_to(canonical_root):
                    unsafe += 1
                    continue
                stat = resolved.stat()
            except (OSError, ValueError):
                unsafe += 1
                continue
            if not resolved.is_file():
                continue
            if resolved.suffix.casefold() not in SUPPORTED_SUFFIXES:
                unsupported += 1
                continue
            if stat.st_size > cap:
                oversized += 1
                errors.append(
                    {
                        "relative_name": path.relative_to(canonical_root).as_posix(),
                        "code": "file_too_large",
                    }
                )
                continue
            relative = resolved.relative_to(canonical_root).as_posix()
            if total_supported_bytes + stat.st_size > settings.KNOWLEDGE_MAX_SCAN_BYTES:
                oversized += 1
                errors.append({"relative_name": relative, "code": "scan_total_limit"})
                continue
            try:
                with resolved.open("rb") as stream:
                    raw = stream.read(cap + 1)
                if len(raw) > cap:
                    oversized += 1
                    errors.append({"relative_name": relative, "code": "file_too_large"})
                    continue
                total_supported_bytes += len(raw)
                file_hash = hashlib.sha256(raw).hexdigest()
                if (
                    known_hashes
                    and relative not in force_parse
                    and known_hashes.get(relative) == file_hash
                ):
                    unchanged.append(
                        (relative, stat.st_size, stat.st_mtime_ns, file_hash)
                    )
                    continue
                parsed = parse_file(relative, raw, stat.st_mtime_ns, cap=cap)
            except DocumentParseError as exc:
                failed += 1
                errors.append({"relative_name": relative, "code": exc.code})
                continue
            except OSError:
                failed += 1
                errors.append({"relative_name": relative, "code": "file_read_failed"})
                continue
            files.append(parsed)
    files.sort(key=lambda item: item.relative_name.casefold())
    return ScanResult(
        files, unsupported, oversized, unsafe, scanned, failed, errors, unchanged
    )


def parse_file(
    relative_name: str, raw: bytes, modified_ns: int, *, cap: int | None = None
) -> ParsedFile:
    if cap is not None and len(raw) > cap:
        raise DocumentParseError("file_too_large")
    suffix = Path(relative_name).suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentParseError("unsupported_file")
    try:
        content, content_type = _extract_text(suffix, raw)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("document_parse_failed") from exc
    if len(content) > MAX_EXTRACTED_TEXT_CHARS:
        raise DocumentParseError("document_text_limit")
    normalized = content.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        raise DocumentParseError("document_empty")
    markdown = parse_markdown(normalized) if suffix == ".md" else None
    digest = hashlib.sha256(raw).hexdigest()
    modified = datetime.fromtimestamp(modified_ns / 1_000_000_000, UTC)
    sensitive = any(pattern.search(normalized) for pattern in _SECRET_PATTERNS)
    return ParsedFile(
        relative_name,
        content_type,
        normalized,
        digest,
        len(raw),
        modified,
        modified_ns,
        markdown,
        sensitive,
    )


def _extract_text(suffix: str, raw: bytes) -> tuple[str, str]:
    if suffix in {".md", ".txt"}:
        try:
            return raw.decode(
                "utf-8-sig"
            ), "text/markdown" if suffix == ".md" else "text/plain"
        except UnicodeDecodeError as exc:
            raise DocumentParseError("document_encoding_unsupported") from exc
    if suffix in {".html", ".htm"}:
        parser = _TextHTMLParser()
        parser.feed(raw.decode("utf-8-sig", errors="replace"))
        parser.close()
        return "\n".join(parser.parts), "text/html"
    if suffix == ".json":
        try:
            value = json.loads(raw.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise DocumentParseError("document_json_invalid") from exc
        return json.dumps(value, ensure_ascii=False, indent=2)[
            :2_000_000
        ], "application/json"
    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(text), strict=True))
        except (UnicodeDecodeError, csv.Error) as exc:
            raise DocumentParseError("document_csv_invalid") from exc
        return "\n".join(" | ".join(cell for cell in row) for row in rows)[
            :2_000_000
        ], "text/csv"
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw), strict=False)
        text = "\n\n".join(
            f"[Page {index + 1}]\n{page.extract_text() or ''}"
            for index, page in enumerate(reader.pages[:500])
        )
        return text, "application/pdf"
    if suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 2_000
                or sum(entry.file_size for entry in entries)
                > MAX_DOCX_UNCOMPRESSED_BYTES
            ):
                raise DocumentParseError("document_archive_limit")
            if any(
                name.lower().endswith("vbaproject.bin") for name in archive.namelist()
            ):
                # Macros are never evaluated; a macro-bearing document is excluded.
                raise DocumentParseError("document_macro_content_rejected")
        from docx import Document

        document = Document(io.BytesIO(raw))
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n\n".join(
            parts
        ), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raise DocumentParseError("unsupported_file")


class _TextHTMLParser(HTMLParser):
    _blocked = {"script", "style", "iframe", "object", "embed", "svg", "noscript"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.block_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.casefold() in self._blocked:
            self.block_depth += 1
        elif not self.block_depth and tag.casefold() in {
            "p",
            "br",
            "div",
            "h1",
            "h2",
            "h3",
            "li",
            "tr",
        }:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.casefold() in self._blocked and self.block_depth:
            self.block_depth -= 1
        elif not self.block_depth and tag.casefold() in {
            "p",
            "div",
            "h1",
            "h2",
            "h3",
            "li",
            "tr",
        }:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.block_depth:
            clean = html.unescape(data).strip()
            if clean:
                self.parts.append(clean)


def _is_link(path: Path) -> bool:
    try:
        if path.is_symlink():
            return True
        junction = getattr(path, "is_junction", None)
        return bool(junction and junction())
    except OSError:
        return True
