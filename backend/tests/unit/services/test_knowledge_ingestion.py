from __future__ import annotations

import hashlib
import io
from datetime import UTC, datetime
from pathlib import Path

import pytest
from app.services.knowledge.ingestion import (
    DocumentParseError,
    parse_file,
    scan_directory,
)
from app.services.knowledge.source_service import (
    _expand_markdown_embeds,
    _safe_relative_name,
    _write_uploaded_files,
)


def _parse(name: str, text: str):
    modified_ns = int(datetime.now(UTC).timestamp() * 1_000_000_000)
    return parse_file(name, text.encode("utf-8"), modified_ns)


def test_safe_parsers_strip_active_html_and_normalize_text() -> None:
    html_doc = _parse(
        "guide.html",
        "<h1>Guide</h1><script>do not index this token</script><p>Safe text</p>",
    )
    json_doc = _parse("reference.json", '{"title":"IR","steps":["contain","recover"]}')
    csv_doc = _parse("table.csv", "name,value\nservice,expected\n")
    txt_doc = _parse("note.txt", "Local incident checklist")

    assert "Safe text" in html_doc.content
    assert "do not index" not in html_doc.content
    assert "contain" in json_doc.content
    assert "service | expected" in csv_doc.content
    assert txt_doc.content_type == "text/plain"


def test_pdf_and_docx_parse_without_executing_active_content() -> None:
    from docx import Document
    from reportlab.pdfgen.canvas import Canvas

    docx_buffer = io.BytesIO()
    document = Document()
    document.add_heading("Response guide", level=1)
    document.add_paragraph("Preserve evidence and document the source.")
    document.save(docx_buffer)

    pdf_buffer = io.BytesIO()
    canvas = Canvas(pdf_buffer)
    canvas.drawString(40, 760, "Incident response reference")
    canvas.save()

    docx_result = parse_file("response.docx", docx_buffer.getvalue(), 1)
    pdf_result = parse_file("response.pdf", pdf_buffer.getvalue(), 1)
    assert "Preserve evidence" in docx_result.content
    assert "Incident response reference" in pdf_result.content
    assert "[Page 1]" in pdf_result.content


def test_invalid_documents_return_safe_classified_errors() -> None:
    with pytest.raises(DocumentParseError, match="document_json_invalid"):
        parse_file("broken.json", b"{invalid", 1)
    with pytest.raises(DocumentParseError, match="document_encoding_unsupported"):
        parse_file("broken.md", b"\xff\xfe", 1)
    with pytest.raises(DocumentParseError, match="file_too_large"):
        parse_file("large.txt", b"x" * 10, 1, cap=5)


def test_secret_warning_detects_risk_without_returning_secret_value() -> None:
    parsed = _parse("sensitive.md", "api_key=ghp_12345678901234567890")
    assert parsed.sensitive_content_warning
    assert "ghp_12345678901234567890" not in parsed.file_hash


def test_selected_source_paths_reject_traversal_and_skip_obsidian_plugin_data() -> None:
    assert _safe_relative_name("nested/Guía de IR.md") == "nested/Guía de IR.md"
    assert _safe_relative_name(".obsidian/plugins/plugin.js") is None
    with pytest.raises(ValueError, match="unsafe"):
        _safe_relative_name("../outside.md")
    with pytest.raises(ValueError, match="unsafe"):
        _safe_relative_name("C:\\outside.md")


async def test_snapshot_upload_rejects_symlinked_parent_before_writing(
    tmp_path: Path,
) -> None:
    root = tmp_path / "RavenTech OSINT" / "knowledge" / "source"
    root.mkdir(parents=True)
    outside = tmp_path / "outside"
    outside.mkdir()
    try:
        (root / "linked").symlink_to(outside, target_is_directory=True)
    except OSError:
        pytest.skip("Symlink creation is unavailable in this test environment.")

    with pytest.raises(ValueError, match="unsafe link"):
        await _write_uploaded_files(root, [("linked/escape.md", b"must not write", 1)])

    assert not (outside / "escape.md").exists()


async def test_snapshot_upload_supports_unicode_and_spaces(tmp_path: Path) -> None:
    root = tmp_path / "RavenTech OSINT" / "knowledge" / "source"
    root.mkdir(parents=True)
    selected = await _write_uploaded_files(
        root,
        [("Guía de respuesta/初動.md", b"contenido local", 1)],
    )
    assert selected == ["Guía de respuesta/初動.md"]
    assert (root / "Guía de respuesta" / "初動.md").read_text(
        encoding="utf-8"
    ) == "contenido local"


def test_scan_rejects_symlinks_and_ignores_vault_noise(tmp_path: Path) -> None:
    root = tmp_path / "vault"
    root.mkdir()
    (root / "note.md").write_text("# Note\n", encoding="utf-8")
    plugin = root / ".obsidian"
    plugin.mkdir()
    (plugin / "config.md").write_text("not indexed", encoding="utf-8")
    external = tmp_path / "outside.md"
    external.write_text("outside root", encoding="utf-8")
    link = root / "escape.md"
    try:
        link.symlink_to(external)
    except OSError:
        pytest.skip("Symlink creation is unavailable in this test environment.")

    result = scan_directory(str(root))
    assert [item.relative_name for item in result.files] == ["note.md"]
    assert result.unsafe_links == 1


def test_scan_hashes_unchanged_files_without_reparsing_and_can_force_embed_dependencies(
    tmp_path: Path,
) -> None:
    root = tmp_path / "vault"
    root.mkdir()
    note = root / "Note.md"
    raw = b"# Note\n\nStable content."
    note.write_bytes(raw)
    digest = hashlib.sha256(raw).hexdigest()

    unchanged = scan_directory(str(root), known_hashes={"Note.md": digest})
    assert unchanged.files == []
    assert unchanged.unchanged[0][0] == "Note.md"
    assert unchanged.unchanged[0][1] == len(raw)
    assert unchanged.unchanged[0][3] == digest

    forced = scan_directory(
        str(root), known_hashes={"Note.md": digest}, force_parse_names={"Note.md"}
    )
    assert [item.relative_name for item in forced.files] == ["Note.md"]
    assert forced.unchanged == []


def test_scan_handles_synthetic_500_note_vault_incrementally(tmp_path: Path) -> None:
    root = tmp_path / "large-vault"
    root.mkdir()
    for index in range(500):
        (root / f"Note {index:04}.md").write_text(
            f"# Note {index}\n\nLocal defensive reference {index}.\n",
            encoding="utf-8",
        )

    first = scan_directory(str(root))
    assert len(first.files) == 500
    previous_hashes = {item.relative_name: item.file_hash for item in first.files}
    (root / "Note 0250.md").write_text(
        "# Note 250\n\nUpdated local defensive reference.\n", encoding="utf-8"
    )

    second = scan_directory(str(root), known_hashes=previous_hashes)
    assert len(second.files) == 1
    assert second.files[0].relative_name == "Note 0250.md"
    assert len(second.unchanged) == 499


def test_embeds_expand_one_level_and_incrementally_reindex() -> None:
    main = _parse("Main.md", "# Main\n\n![[Appendix]]")
    appendix = _parse(
        "Appendix.md", "# Appendix\n\nEmbedded control evidence\n\n![[Deep]]"
    )
    deep = _parse("Deep.md", "# Deep\n\nDeep content is not recursively embedded")
    expanded = _expand_markdown_embeds([main, appendix, deep])
    main_result = next(item for item in expanded if item.relative_name == "Main.md")
    appendix_result = next(
        item for item in expanded if item.relative_name == "Appendix.md"
    )
    assert "Embedded control evidence" in main_result.content
    assert main_result.embed_hashes == {"Appendix.md": appendix.file_hash}
    assert "Deep content is not recursively embedded" not in main_result.content
    assert "Deep content is not recursively embedded" in appendix_result.content


def test_parse_timestamp_is_utc_aware() -> None:
    parsed = _parse("note.txt", "note")
    assert parsed.modified_at.tzinfo == UTC
